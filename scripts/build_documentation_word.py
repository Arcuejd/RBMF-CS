"""Read five requested project sources and compile a Word reference manual."""
from pathlib import Path
import hashlib
import json
import re
from docx import Document
from docx.enum.section import WD_SECTION_START, WD_ORIENT
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT/'outputs/documentation_word'
SOURCES = ['README.md', 'docs/math_to_code.md', 'configs/synthetic_demo.json',
           'docs/validation_report.md', 'docs/stage3_validation.md']
HEADINGS = ['一 项目说明与运行指南', '二 公式映射 变量维度与歧义记录',
            '三 Synthetic Demo 统一配置', '四 阶段 1 和 2 验证报告', '五 阶段 3 完整验证报告']
doc = Document()
ledger = []
# Remove template paragraph borders, including the inherited blue Title rule.
for style in doc.styles:
    for element in list(style.element.iter(qn('w:pBdr'))):
        element.getparent().remove(element)


def font(style, name, east_asia, size):
    style.font.name = name
    style.font.size = Pt(size)
    style.font.color.rgb = RGBColor(0, 0, 0)
    properties = style.element.get_or_add_rPr()
    fonts = properties.find(qn('w:rFonts'))
    if fonts is None:
        fonts = OxmlElement('w:rFonts'); properties.append(fonts)
    for key in ('ascii', 'hAnsi', 'cs'):
        fonts.set(qn('w:'+key), name)
    fonts.set(qn('w:eastAsia'), east_asia)


for name in ('Normal','Body Text','List Paragraph'):
    font(doc.styles[name], 'Calibri', 'Microsoft YaHei', 10.5)
    pf = doc.styles[name].paragraph_format
    pf.space_after = Pt(6); pf.line_spacing = 1.18
    pf.widow_control = True
for name, size in [('Title',24),('Subtitle',12),('Heading 1',18),('Heading 2',14),('Heading 3',11.5)]:
    font(doc.styles[name], 'Calibri', 'Microsoft YaHei', size)
    doc.styles[name].paragraph_format.space_before = Pt(12)
    doc.styles[name].paragraph_format.space_after = Pt(8)
    doc.styles[name].paragraph_format.keep_with_next = True
code_style = doc.styles.add_style('Source Code', 1)
font(code_style, 'Consolas', 'Microsoft YaHei', 8.5)
code_style.paragraph_format.line_spacing = 1.05
code_style.paragraph_format.space_after = Pt(0)
code_style.paragraph_format.left_indent = Cm(.25)
code_style.paragraph_format.keep_with_next = False


def section_layout(section, landscape=False):
    section.orientation = WD_ORIENT.LANDSCAPE if landscape else WD_ORIENT.PORTRAIT
    section.page_width, section.page_height = (Cm(29.7), Cm(21)) if landscape else (Cm(21), Cm(29.7))
    section.top_margin = section.bottom_margin = Cm(1.8)
    section.left_margin = section.right_margin = Cm(1.7)
    section.footer_distance = Cm(.8)
    footer = section.footer.paragraphs[0]
    if not footer.runs:
        footer.alignment = 2
        run = footer.add_run()
        field = OxmlElement('w:fldSimple'); field.set(qn('w:instr'), 'PAGE')
        run._r.addnext(field)
    # Page numbering and table continuation make this long manual navigable.


section_layout(doc.sections[0])
doc.core_properties.title = 'RBMF CS 科研项目说明与验证汇编'
doc.core_properties.subject = '运行说明 公式映射 配置与阶段验证记录'
doc.core_properties.author = 'RBMF-CS Project'
doc.add_paragraph('RBMF CS 科研项目说明与验证汇编', 'Title')
doc.add_paragraph('Robust Bayesian Multi Fidelity Compressed Sensing', 'Subtitle')
doc.add_paragraph('本册包含项目运行说明、公式与代码对应关系、synthetic demo 配置，以及阶段 1 至 3 的验证记录。各章保留原文件的结果口径和历史日期，便于查阅实现、运行项目和追踪验证证据。')
doc.add_paragraph('整理日期  2026 年 10 月 14 日')
doc.add_paragraph('内容导航', 'Heading 2')
for title, source in zip(HEADINGS, SOURCES):
    p=doc.add_paragraph(title)
    p.add_run('\n'+source).font.size=Pt(9)


def inline(paragraph, text):
    """Preserve source text and code; remove only Markdown presentation markup."""
    pattern = r'(`[^`]+`|\*\*[^*]+\*\*|\[[^\]]+\]\([^\)]+\))'
    for token in re.split(pattern, text):
        if not token:
            continue
        if token.startswith('`') and token.endswith('`'):
            run=paragraph.add_run(token[1:-1])
            run.font.name='Consolas'; run.font.size=Pt(9)
            run._r.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'Microsoft YaHei')
        elif token.startswith('**') and token.endswith('**'):
            paragraph.add_run(token[2:-2]).bold=True
        elif re.fullmatch(r'\[[^\]]+\]\([^\)]+\)',token):
            match=re.fullmatch(r'\[([^\]]+)\]\(([^\)]+)\)',token)
            paragraph.add_run(match[1]+' ('+match[2]+')')
        else:
            paragraph.add_run(token)


def table(rows, landscape, source):
    columns=len(rows[0])
    width=26.3 if landscape else 17.6
    if landscape and columns==4:
        widths=[2.1,9.3,6.8,8.1]
    elif landscape and columns==3:
        widths=[2.2,15.,9.1]
    elif columns==2:
        widths=[width*.36,width*.64]
    elif columns==7:
        widths=[2.,3.,1.6,2.1,1.6,3.8,3.5]
    elif columns==6:
        widths=[1.7,1.6,2.8,3.4,2.,6.1]
    else:
        widths=[width/columns]*columns
    if abs(sum(widths)-width)>.01:
        widths=[x*width/sum(widths) for x in widths]
    t=doc.add_table(rows=0,cols=columns)
    t.alignment=WD_TABLE_ALIGNMENT.CENTER; t.autofit=False
    for col,w in zip(t.columns,widths): col.width=Cm(w)
    pr=t._tbl.tblPr
    borders=OxmlElement('w:tblBorders')
    for key in ('top','left','bottom','right','insideH','insideV'):
        element=OxmlElement('w:'+key)
        for attr,value in [('val','single'),('sz','4'),('color','D9D9D9')]: element.set(qn('w:'+attr),value)
        borders.append(element)
    pr.append(borders)
    for ri,data in enumerate(rows):
        cells=t.add_row().cells
        if ri==0:
            repeat=OxmlElement('w:tblHeader'); t.rows[-1]._tr.get_or_add_trPr().append(repeat)
        # Prevent splitting one semantic mapping row across pages.
        cant=OxmlElement('w:cantSplit'); t.rows[-1]._tr.get_or_add_trPr().append(cant)
        for ci,(cell,text) in enumerate(zip(cells,data)):
            cell.width=Cm(widths[ci]); cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            margins=OxmlElement('w:tcMar')
            for key,value in [('top','70'),('bottom','70'),('left','90'),('right','90')]:
                el=OxmlElement('w:'+key); el.set(qn('w:w'),value); el.set(qn('w:type'),'dxa'); margins.append(el)
            cell._tc.get_or_add_tcPr().append(margins)
            shade=OxmlElement('w:shd'); shade.set(qn('w:fill'),'E7EDF3' if ri==0 else ('F7F9FB' if ri%2==0 else 'FFFFFF'))
            cell._tc.get_or_add_tcPr().append(shade)
            p=cell.paragraphs[0]
            p.paragraph_format.space_after=Pt(2)
            p.paragraph_format.line_spacing=1.1
            if ri==0:
                p.paragraph_format.keep_with_next=True
            inline(p,text)
            for run in p.runs:
                run.font.size=Pt(9 if landscape else 9.5)
                if ri==0: run.bold=True
    ledger.append({'source':source,'kind':'table','rows':len(rows),'columns':columns})
    doc.add_paragraph().paragraph_format.space_after=Pt(2)


def parse_markdown(text, landscape, source):
    lines=text.splitlines(); i=0
    while i<len(lines):
        line=lines[i]
        if not line.strip(): i+=1; continue
        if line.startswith('```'):
            i+=1; code=[]
            while i<len(lines) and not lines[i].startswith('```'):
                code.append(lines[i]); i+=1
            for code_line in code:
                p=doc.add_paragraph(style='Source Code'); p.add_run(code_line)
            doc.add_paragraph().paragraph_format.space_after=Pt(3)
            ledger.append({'source':source,'kind':'code','lines':len(code)})
            i+=1; continue
        if line.startswith('|'):
            rows=[]
            while i<len(lines) and lines[i].startswith('|'):
                raw=lines[i].strip().strip('|').strip()
                if re.fullmatch(r'[\s|:\-]+', raw):
                    i+=1
                    continue
                cells=re.split(r'\s+\|\s+',raw)
                if not all(re.fullmatch(r':?-+:?',x.strip()) for x in cells):
                    rows.append([x.strip() for x in cells])
                i+=1
            if any(len(r)!=len(rows[0]) for r in rows):
                raise ValueError(f'table parse mismatch: {source}: {rows}')
            table(rows,landscape,source); continue
        heading=re.match(r'^(#{1,3})\s+(.*)$',line)
        if heading:
            # Keep the source H1 as a subtitle beneath the chapter navigation.
            level=min(len(heading[1])+1,3)
            p=doc.add_paragraph(style='Heading '+str(level))
            inline(p,heading[2]); i+=1; continue
        bullet=re.match(r'^(- |\d+\. )(.*)$',line)
        if bullet:
            p=doc.add_paragraph(style='List Paragraph')
            p.paragraph_format.left_indent=Cm(.45)
            p.paragraph_format.first_line_indent=Cm(-.35)
            inline(p,('• ' if bullet[1]=='- ' else bullet[1])+bullet[2])
        else:
            p=doc.add_paragraph(); inline(p,line)
        ledger.append({'source':source,'kind':'paragraph'})
        i+=1


for chapter,(source,heading) in enumerate(zip(SOURCES,HEADINGS)):
    section=doc.add_section(WD_SECTION_START.NEW_PAGE)
    landscape=chapter==1
    section_layout(section,landscape)
    doc.add_paragraph(heading,'Heading 1')
    doc.add_paragraph('来源文件  '+source).runs[0].font.size=Pt(9)
    text=(ROOT/source).read_text(encoding='utf-8-sig')
    if source.endswith('.json'):
        config=json.loads(text)
        doc.add_paragraph('配置参数表','Heading 2')
        for group,value in config.items():
            if isinstance(value,dict):
                doc.add_paragraph(group,'Heading 3')
                rows=[['参数','值']]+[[key,json.dumps(val,ensure_ascii=False)] for key,val in value.items()]
                table(rows,False,source)
            else:
                p=doc.add_paragraph(); inline(p,group+' = '+json.dumps(value,ensure_ascii=False))
        doc.add_paragraph('原始 JSON 配置','Heading 2')
        for line in text.splitlines():
            doc.add_paragraph(line,'Source Code')
    else:
        parse_markdown(text,landscape,source)

OUT.mkdir(parents=True,exist_ok=True)
target=OUT/'RBMF_CS_Project_Documentation_20261014.docx'
doc.save(target)
manifest={'output':str(target),'sources':[{'path':str(ROOT/s),'sha256':hashlib.sha256((ROOT/s).read_bytes()).hexdigest()} for s in SOURCES],
          'chapters':HEADINGS,'content_blocks':ledger}
(OUT/'compilation_manifest.json').write_text(json.dumps(manifest,indent=2,ensure_ascii=False),encoding='utf-8')
print(json.dumps({'output':str(target),'tables':len(doc.tables),'paragraphs':len(doc.paragraphs)},ensure_ascii=False))
